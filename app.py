"""KESCO Auto-Coder - local web app.  Run:  python app.py   then open http://127.0.0.1:5000"""
import os, re, io, uuid, difflib, json, webbrowser, threading, unicodedata, collections
from pathlib import Path
import numpy as np, openpyxl, pandas as pd
from flask import Flask, request, jsonify, send_file, Response
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity
from sklearn.preprocessing import normalize

BASE = Path(__file__).parent
REF_SAVE = BASE / "kesco_reference.xlsx"

def _fail(msg): raise ValueError(msg)

# ---------------- matcher (identical logic to kesco_matcher.ipynb) ----------------
CODE_PATTERN = re.compile(r"^\d{3,4}-\d{1,3}$")

# strips gender markers (m/f/d), hours/schedule/company parentheticals, and
# "at <Company> GmbH/Ltd..." tails before matching -- doesn't touch your data
_GENDER_MARK = re.compile(r"\((?:[mwfdx]\s*/\s*){1,3}[mwfdx]\)", re.IGNORECASE)
_NOISE_PAREN = re.compile(
    r"\((?:[^()]*?(?:\d+\s*[-/]?\s*\d*\s*h(?:ours?|rs?)?\b|h/?week|std\.?/?wo\.?|"
    r"full[-\s]?time|part[-\s]?time|vollzeit|teilzeit|all genders|diverse|"
    r"gmbh|ltd\.?|llc|kg\b|\bag\b|inc\.?|co\.?\b)[^()]*?)\)",
    re.IGNORECASE,
)
_COMPANY_TAIL = re.compile(
    r"\b(?:at|@)\s+[A-Z][\w&.,'-]*(?:\s+[A-Z][\w&.,'-]*){0,4}\s*"
    r"(?:GmbH|Ltd\.?|LLC|Inc\.?|KG|AG|Co\.?|Group)\b.*$"
)


def load_kesco_reference(xlsx_path: str) -> pd.DataFrame:
    """Flatten the hierarchical KESCO workbook into {code, title, group ancestry} rows."""
    ws = openpyxl.load_workbook(xlsx_path, data_only=True).active
    norm = lambda g: re.sub(r"\s+", " ", str(g)).strip().lower() if g else None
    major = submajor = minor = unit = None
    records = []

    for g_raw, code, desc in (row[:3] for row in ws.iter_rows(min_row=2, values_only=True)):
        g, desc = norm(g_raw), (str(desc).strip() if desc else None)
        code_str = str(code).strip() if code is not None else None

        if g and code_str and not CODE_PATTERN.match(code_str):
            if g.startswith("major") and "sub" not in g: major = desc
            elif g.startswith("sub"): submajor = desc
            elif g.startswith("minor"): minor = desc
            elif g.startswith("unit"): unit = desc
            continue

        if code_str and CODE_PATTERN.match(code_str) and desc:
            records.append(dict(code=code_str, title=desc, major_group=major,
                                 sub_major_group=submajor, minor_group=minor, unit_group=unit))

    df = pd.DataFrame(records).drop_duplicates(subset=["code"]).reset_index(drop=True)
    if df.empty:
        raise ValueError("No occupation codes found -- expected columns [GROUP, CODE, DESCRIPTION].")
    return df


def _clean(text: str) -> str:
    text = re.sub(r"[\u200b\u200c\u200d\ufeff]", "", str(text))
    text = re.split(r"\s*\|\s*", text)[0]
    text = re.split(r"\bLocation\s*:", text, flags=re.IGNORECASE)[0]
    text = re.split(r"\bwanted\.", text, flags=re.IGNORECASE)[0]
    text = _COMPANY_TAIL.sub("", text)
    text = _NOISE_PAREN.sub(" ", text)
    text = _GENDER_MARK.sub(" ", text.lower())
    text = re.sub(r"[^a-z0-9\s]", " ", text)
    return re.sub(r"\s+", " ", text).strip()


class KESCOMatcher:
    def __init__(self, reference, confidence_threshold: float = 0.45):
        """`reference`: path to KESCO .xlsx/.json/.csv, or an already-parsed DataFrame."""
        self.threshold = confidence_threshold
        if isinstance(reference, pd.DataFrame):
            self.ref = reference
        else:
            suf = Path(reference).suffix.lower()
            self.ref = (load_kesco_reference(reference) if suf in (".xlsx", ".xls") else
                        pd.read_json(reference) if suf == ".json" else
                        pd.read_csv(reference, dtype=str) if suf == ".csv" else
                        _fail("reference must be .xlsx, .json, .csv, or a DataFrame"))

        self._clean_titles = self.ref["title"].map(_clean).tolist()
        self.word_vec = TfidfVectorizer(ngram_range=(1, 2)).fit(self._clean_titles)
        self.char_vec = TfidfVectorizer(analyzer="char_wb", ngram_range=(3, 5)).fit(self._clean_titles)
        self.word_matrix = self.word_vec.transform(self._clean_titles)
        self.char_matrix = self.char_vec.transform(self._clean_titles)

    def match(self, title: str, top_n: int = 3):
        """Top-N candidate KESCO matches for one free-text title."""
        return self._match_many([title], top_n)[0]

    def _match_many(self, titles, top_n=3, prefilter_k=25, chunk_size=500):
        """Vectorised matching (fast on large datasets): TF-IDF narrows candidates,
        difflib only scores each row's top `prefilter_k` candidates."""
        clean_list = [_clean(t) for t in titles]
        results = [None] * len(clean_list)

        for start in range(0, len(clean_list), chunk_size):
            chunk = clean_list[start:start + chunk_size]
            word_sim = (self.word_vec.transform(chunk) @ self.word_matrix.T).toarray()
            char_sim = (self.char_vec.transform(chunk) @ self.char_matrix.T).toarray()
            pre_blend = 0.45 * word_sim + 0.35 * char_sim

            for i, clean in enumerate(chunk):
                if not clean:
                    results[start + i] = []
                    continue
                k = min(prefilter_k, pre_blend.shape[1])
                cand = np.argpartition(pre_blend[i], -k)[-k:]
                scored = sorted(
                    ((idx, 0.45 * word_sim[i, idx] + 0.35 * char_sim[i, idx]
                      + 0.20 * difflib.SequenceMatcher(None, clean, self._clean_titles[idx]).ratio())
                     for idx in cand),
                    key=lambda x: x[1], reverse=True,
                )[:top_n]
                results[start + i] = [self._row_dict(idx, s) for idx, s in scored]
        return results

    def _row_dict(self, idx, score):
        r = self.ref.iloc[idx]
        return {"code": r.code, "title": r.title, "confidence": round(float(score), 4),
                "major_group": r.major_group, "sub_major_group": r.sub_major_group,
                "minor_group": r.minor_group, "unit_group": r.unit_group}

    def match_batch(self, titles, top_n=3) -> pd.DataFrame:
        """Match a plain list of titles -> standalone lookup DataFrame."""
        rows = []
        for t, matches in zip(titles, self._match_many(list(titles), top_n)):
            if not matches:
                rows.append({"input_title": t, "status": "NEEDS_REVIEW"}); continue
            best = matches[0]
            row = {"input_title": t, "matched_code": best["code"], "matched_title": best["title"],
                   "confidence": best["confidence"], **{k: best[k] for k in
                   ("major_group", "sub_major_group", "minor_group", "unit_group")},
                   "status": "MATCHED" if best["confidence"] >= self.threshold else "NEEDS_REVIEW"}
            for j, alt in enumerate(matches[1:top_n], 1):
                row[f"alt{j}_code"], row[f"alt{j}_title"], row[f"alt{j}_confidence"] = \
                    alt["code"], alt["title"], alt["confidence"]
            rows.append(row)
        return pd.DataFrame(rows)

    def map_dataset(self, df: pd.DataFrame, occupation_column: str, code_column="kesco_code") -> pd.DataFrame:
        """Return a COPY of df with kesco_code* columns appended (original columns untouched).
        Matches unique titles once, then maps results back onto every row."""
        if occupation_column not in df.columns:
            raise KeyError(f"'{occupation_column}' not in columns: {list(df.columns)}")
        out = df.copy()
        titles = out[occupation_column].astype(str)
        uniq = titles.unique().tolist()
        by_title = dict(zip(uniq, self._match_many(uniq, top_n=1)))

        fields = ["code", "title", "confidence", "major_group", "sub_major_group", "minor_group", "unit_group"]
        cols = {f: [] for f in fields}
        statuses = []
        for t in titles:
            m = by_title.get(t) or [{}]
            best = m[0]
            for f in fields:
                cols[f].append(best.get(f))
            statuses.append("MATCHED" if (best.get("confidence") or 0) >= self.threshold else "NEEDS_REVIEW")

        out[code_column] = cols["code"]
        out[f"{code_column}_title"] = cols["title"]
        out[f"{code_column}_confidence"] = cols["confidence"]
        out[f"{code_column}_status"] = statuses
        out[f"{code_column}_major_group"] = cols["major_group"]
        out[f"{code_column}_sub_major_group"] = cols["sub_major_group"]
        out[f"{code_column}_minor_group"] = cols["minor_group"]
        out[f"{code_column}_unit_group"] = cols["unit_group"]
        return out

    def map_file(self, input_path, occupation_column, output_path=None, sheet_name=0, code_column="kesco_code"):
        """Read a dataset file, append KESCO columns, write it back out."""
        p = Path(input_path)
        df = (pd.read_excel(input_path, sheet_name=sheet_name) if p.suffix.lower() in (".xlsx", ".xls")
              else pd.read_csv(input_path) if p.suffix.lower() == ".csv"
              else _fail(f"Unsupported file type: {p.suffix}"))
        mapped = self.map_dataset(df, occupation_column, code_column=code_column)

        out_p = Path(output_path or p.with_name(p.stem + "_KESCO_mapped" + p.suffix))
        out_p.parent.mkdir(parents=True, exist_ok=True)
        (mapped.to_excel if out_p.suffix.lower() in (".xlsx", ".xls") else mapped.to_csv)(out_p, index=False)
        return mapped

    def map_multiple_files(self, file_specs, output_dir=None, code_column="kesco_code"):
        """Batch-process several country/county datasets, each with its own column name.
        file_specs: [{"path": ..., "occupation_column": ..., "sheet_name": 0 (optional)}, ...]"""
        results = {}
        for spec in file_specs:
            p = Path(spec["path"])
            out_path = str(Path(output_dir) / (p.stem + "_KESCO_mapped" + p.suffix)) if output_dir else spec.get("output_path")
            mapped = self.map_file(spec["path"], spec["occupation_column"], out_path,
                                    spec.get("sheet_name", 0), code_column)
            results[spec["path"]] = mapped
            n_review = (mapped[f"{code_column}_status"] == "NEEDS_REVIEW").sum()
            print(f"{spec['path']}: {len(mapped)} rows, {n_review} flagged NEEDS_REVIEW")
        return results

PAGE_CSS = r"""
:root{--bg:#f6f7f9;--card:#fff;--ink:#1c2430;--mut:#667085;--acc:#0a56b3;--warn:#b54708;--line:#e4e7ec}
*{box-sizing:border-box}body{margin:0;font:15px system-ui,Segoe UI,sans-serif;background:var(--bg);color:var(--ink)}
header{padding:22px 32px;background:linear-gradient(135deg,#083f86,var(--acc));color:#fff}header h1{margin:0;font-size:22px}header p{margin:4px 0 0;opacity:.85}
main{max-width:1250px;margin:24px auto;padding:0 20px}.card{background:var(--card);border:1px solid var(--line);border-radius:12px;padding:20px;margin-bottom:18px}
#drop{border:2px dashed #98a2b3;border-radius:12px;padding:38px;text-align:center;cursor:pointer;color:var(--mut)}#drop.on{border-color:var(--acc);background:#eaf2ff}
button,.btn,select,input[type=text]{font:inherit;padding:8px 14px;border-radius:8px;border:1px solid var(--line);background:#fff;cursor:pointer}
.btn.p,button.p{background:var(--acc);color:#fff;border-color:var(--acc)}.row{display:flex;gap:12px;align-items:center;flex-wrap:wrap}
.stat{padding:10px 16px;border-radius:10px;background:var(--bg)}.stat b{font-size:20px;display:block}
.wrap{overflow:auto;max-height:560px;border:1px solid var(--line);border-radius:8px}table{border-collapse:collapse;width:100%;font-size:13px}
th,td{padding:7px 10px;border-bottom:1px solid var(--line);white-space:nowrap;text-align:left;max-width:320px;overflow:hidden;text-overflow:ellipsis}
th{position:sticky;top:0;background:#f2f4f7}th.k{background:#cfe2ff}td.k{background:#f3f8ff}tr.r td.k{background:#fef0c7}
.tag{padding:2px 8px;border-radius:99px;font-size:12px}.M{background:#cfe2ff;color:#083f86}.R{background:#fef0c7;color:var(--warn)}
#msg{color:#b42318;margin-top:8px}.hide{display:none}"""

# ---------------- web app: KeSCO + KeSic ----------------
JSON_REF = BASE / "kesco_kesic_reference.json"
app = Flask(__name__)
PUBLIC = os.environ.get("PUBLIC") == "1"
app.config["MAX_CONTENT_LENGTH"] = (25 if PUBLIC else 100) * 1024 * 1024
STATE = {"ns": None, "old": None, "ref": None}
JOBS = {}
fmt = lambda c: f"{c[:4]}-{c[4:]}" if c else ""
TGUESS = re.compile(r"job.?title|occupation|position|designation|title|role", re.I)
CGUESS = re.compile(r"company|employer|organi[sz]ation|firm", re.I)

def load_coder():
    ref = json.load(open(JSON_REF, encoding="utf-8"))
    ns = dict(re=re, unicodedata=unicodedata, collections=collections, TfidfVectorizer=TfidfVectorizer, normalize=normalize,
              KESCO_TITLES=ref["kesco_titles"], KESIC_DESC=ref["kesic_desc"], TITLE_DEC=ref["title_decisions"],
              COMPANY_DEC=ref["company_decisions"], UNNAMED={u.lower() for u in ref["unnamed_companies"]},
              RAW_DEC=ref.setdefault("raw_title_decisions", {}))
    exec(compile((BASE / "coder_core.py").read_text(encoding="utf-8"), "coder_core.py", "exec"), ns)
    try: df = load_kesco_reference(str(REF_SAVE))
    except Exception:  # fall back to the official titles inside the JSON
        df = pd.DataFrame([dict(code=fmt(c), title=t, major_group=None, sub_major_group=None, minor_group=None, unit_group=None)
                           for c, t in ref["kesco_titles"].items()])
    STATE.update(ref=ref, ns=ns, old=KESCOMatcher(df, 0.45))

try: load_coder()
except Exception as e: print("Reference not loaded:", e)

def read_any(f, **kw):
    return pd.read_csv(f, **kw) if f.filename.lower().endswith(".csv") else pd.read_excel(f, **kw)

def guess(cols, rx, default=None):
    for c in cols:
        if rx.search(str(c)): return str(c)
    return default

NEWCOLS = {"kesco code", "kesco title", "kesco status", "kesic code", "kesic title", "kesic status"}

def code_dataset(df, tcol, ccol):
    ns, old = STATE["ns"], STATE["old"]
    base = df.drop(columns=[c for c in df.columns if str(c).strip().lower() in NEWCOLS])
    titles = base[tcol].fillna("").astype(str).tolist()
    uniq = list(dict.fromkeys(titles))
    new = ns["code_titles"](uniq)
    need = [t for t in uniq if new[t][1] not in ("exact", "close match")]   # second opinion only where needed
    second = dict(zip(need, old._match_many(need, top_n=1, chunk_size=100))) if need else {}
    res = {}
    for t in uniq:
        code, status, via = new[t]
        o = (second.get(t) or [{}])[0]
        oc = (o.get("code") or "").replace("-", ""); conf = o.get("confidence") or 0
        if status not in ("exact", "close match") and oc:
            if code == oc and conf >= old.threshold: status = "likely (both methods agree)"
            elif conf >= old.threshold: code, status = oc, f"REVIEW (methods disagree, {conf:.2f})"
        res[t] = (code, status)
    out = base.copy()
    pos = list(out.columns).index(tcol) + 1
    out.insert(pos, "kesco code", [fmt(res[t][0]) for t in titles])
    out.insert(pos + 1, "kesco title", [ns["KESCO_TITLES"].get(res[t][0], "") for t in titles])
    out.insert(pos + 2, "kesco status", [res[t][1] for t in titles])
    meta = pd.DataFrame(index=out.index)
    meta["needs_review"] = [s.startswith("REVIEW") for s in out["kesco status"]]
    meta["kesic suggestion"] = ""
    if ccol:
        cos = base[ccol].fillna("").astype(str).tolist()
        cc = {c: ns["code_company"](c) for c in dict.fromkeys(cos)}
        rev = [cc[c][1].startswith("REVIEW") for c in cos]
        codes = ["" if r else (cc[c][0] or "") for c, r in zip(cos, rev)]
        p2 = list(out.columns).index(ccol) + 1
        out.insert(p2, "kesic code", codes)
        out.insert(p2 + 1, "kesic title", [ns["KESIC_DESC"].get(k, "") for k in codes])
        out.insert(p2 + 2, "kesic status", ["REVIEW" if r else cc[c][1].split(" (")[0] for c, r in zip(cos, rev)])
        meta["kesic suggestion"] = [(cc[c][0] or "") if r else "" for c, r in zip(cos, rev)]
        meta["needs_review"] = meta["needs_review"] | pd.Series(rev, index=out.index)
    return out, meta

RUN_LOCK = threading.Lock()   # one coding job at a time keeps memory low

def run(job, tcol, ccol):
    job["tcol"], job["ccol"] = tcol, ccol
    with RUN_LOCK:
        m, meta = code_dataset(job["raw"], tcol, ccol)
        import gc; gc.collect()
    job["mapped"], job["meta"] = m, meta
    st = m["kesco status"].str.split(" ", n=1).str[0].value_counts()
    view = m.head(2000).copy(); view["needs_review"] = meta["needs_review"].head(2000)
    return {"id": job["id"], "name": job["name"], "tcol": tcol, "ccol": ccol or "", "columns": [str(c) for c in job["raw"].columns if str(c).strip().lower() not in NEWCOLS],
            "rows": len(m), "exact": int(st.get("exact", 0)), "close": int(st.get("close", 0)),
            "likely": int(st.get("likely", 0)), "review": int(meta["needs_review"].sum()),
            "kesic": int((m["kesic code"] != "").sum()) if ccol else 0,
            "cols": [str(c) for c in view.columns], "data": json.loads(view.to_json(orient="records", date_format="iso"))}

@app.get("/")
def index(): return Response(PAGE, mimetype="text/html")

@app.get("/api/status")
def status(): return jsonify(ready=STATE["ns"] is not None, public=PUBLIC)

@app.post("/api/upload")
def upload():
    if not STATE["ns"]: return jsonify(error="kesco_kesic_reference.json is missing on the server."), 400
    try: df = read_any(request.files["file"])
    except Exception as e: return jsonify(error=f"Could not read file: {e}"), 400
    cols = list(df.columns)
    tcol = guess(cols, TGUESS, str(cols[0])); ccol = guess([c for c in cols if str(c) != tcol], CGUESS)
    while len(JOBS) >= 4: JOBS.pop(next(iter(JOBS)))
    jid = uuid.uuid4().hex[:10]
    job = JOBS[jid] = {"id": jid, "name": request.files["file"].filename, "raw": df}
    return jsonify(run(job, tcol, ccol))

@app.post("/api/remap")
def remap():
    d = request.json
    return jsonify(run(JOBS[d["id"]], d["tcol"], d.get("ccol") or None))

@app.post("/api/learn")
def learn():
    if PUBLIC: return jsonify(error="Learning is disabled on the public site."), 403
    df = read_any(request.files["file"], dtype=str); tc, cc = request.form["tcol"], request.form.get("ccol")
    ref, ns = STATE["ref"], STATE["ns"]; at = ac = 0
    for _, r in df.iterrows():
        t = r.get(tc); ko = str(r.get("kesco code", "")).replace("-", "").strip()
        if pd.notna(t) and ko in ref["kesco_titles"]:
            keys = ns["title_keys"](t)
            if keys and ref["title_decisions"].get(keys[0]) != ko: ref["title_decisions"][keys[0]] = ko; at += 1
            elif not keys: ref["raw_title_decisions"][str(t).strip()] = ko; at += 1
        if cc:
            comp = ns["clean_company"](r.get(cc)); ki = str(r.get("kesic code", "")).strip()
            if comp and comp.lower() not in ns["UNNAMED"] and ki in ref["kesic_desc"] and ref["company_decisions"].get(comp) != ki:
                ref["company_decisions"][comp] = ki; ac += 1
    json.dump(ref, open(JSON_REF, "w", encoding="utf-8"), ensure_ascii=False); load_coder()
    return jsonify(titles=at, companies=ac)

@app.get("/download/<jid>")
def download(jid):
    if jid not in JOBS: return Response("Session expired - please upload the file again.", 410)
    job = JOBS[jid]; m = job["mapped"]; meta = job["meta"]; stem = Path(job["name"]).stem; buf = io.BytesIO()
    if request.args.get("fmt") == "csv":
        buf.write(m.to_csv(index=False).encode("utf-8-sig")); buf.seek(0)
        return send_file(buf, as_attachment=True, download_name=f"{stem}_CODED.csv", mimetype="text/csv")
    rv = m[meta["needs_review"]].copy(); rv.insert(0, "row", rv.index + 2)
    keep = ["row", job["tcol"], "kesco code", "kesco title", "kesco status"]
    if job["ccol"]:
        rv["kesic suggestion"] = meta.loc[rv.index, "kesic suggestion"]
        keep += [job["ccol"], "kesic code", "kesic title", "kesic status", "kesic suggestion"]
    from openpyxl.styles import PatternFill
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        m.to_excel(w, sheet_name="Coded", index=False)
        rv[keep].to_excel(w, sheet_name="Coding review", index=False)
        ws = w.sheets["Coded"]; ci = list(m.columns).index("kesco code") + 1
        for i, flag in enumerate(meta["needs_review"], start=2):
            if flag: ws.cell(i, ci).fill = PatternFill("solid", fgColor="FFF2CC")
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=f"{stem}_CODED.xlsx")

PAGE = r"""<!doctype html><html><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>KeSCO &amp; KeSic Auto-Coder - Kenya occupation and industry coding tool</title>
<meta name="description" content="Free online tool: upload an Excel or CSV of job vacancies and automatically code job titles to KeSCO and companies to KeSic.">
<style>CSSHERE
.big{font-size:17px;padding:12px 34px;font-weight:600}#drop{display:flex;flex-direction:column;align-items:center;gap:12px}
#bar{height:8px;background:#e4e7ec;border-radius:99px;overflow:hidden;margin-top:12px}#barfill{display:block;height:100%;width:0;background:var(--acc);transition:width .2s}
#barfill.busy{width:40%!important;animation:slide 1.3s infinite ease-in-out}@keyframes slide{0%{margin-left:-40%}100%{margin-left:100%}}
#stmsg{font-weight:600;color:var(--acc)}
</style></head><body>
<header><h1>KLMIS &middot; KeSCO &amp; KeSic Auto-Coder</h1><p>Upload a vacancies Excel/CSV &rarr; job titles are coded to KeSCO and companies to KeSic automatically.</p></header><main>
<div class="card" id="refbox"><span id="reftxt">Checking reference...</span></div>
<div class="card"><div id="drop"><button class="p big" id="pick" type="button">Upload file</button><div id="dtxt">or drop your Excel / CSV file here</div><input type="file" id="file" accept=".xlsx,.xls,.csv" hidden></div>
<div id="stage" class="hide"><div id="stmsg"></div><div id="bar"><i id="barfill"></i></div></div><div id="msg"></div></div>
<div id="out" class="hide"><div class="card"><div class="row"><div class="stat"><b id="sRows">0</b>rows</div><div class="stat"><b id="sE">0</b>exact</div>
<div class="stat"><b id="sC">0</b>close / likely</div><div class="stat"><b id="sR">0</b>need review</div><div class="stat"><b id="sK">0</b>KeSic coded</div></div>
<div class="row" style="margin-top:14px"><span>Job title column:</span><select id="tcol"></select><span>Company column:</span><select id="ccol"></select><button id="remap">Re-code</button></div>
<div class="row" style="margin-top:14px"><button class="p" id="viewBtn">View results</button><a class="btn p" id="dx">Download Excel</a><a class="btn" id="dc">Download CSV</a>
<label class="btn" id="learnLbl">Upload corrected file (teach the coder)<input type="file" id="learn" accept=".xlsx,.csv" hidden></label></div></div>
<div class="card hide" id="view"><div class="row" style="margin-bottom:10px"><select id="flt"><option value="">All rows</option><option value="r">Needs review only</option><option value="ok">Trusted only</option></select>
<input type="text" id="q" placeholder="Search..." style="flex:1"><span id="cnt" style="color:var(--mut)"></span></div><div class="wrap"><table id="t"></table></div></div></div></main>
<script>
const $=id=>document.getElementById(id);let cur=null,pub=false;
fetch('/api/status').then(r=>r.json()).then(r=>{pub=r.public;$('reftxt').innerHTML=r.ready?'KeSCO / KeSic reference loaded (built in).':'<b style="color:#b54708">Reference file missing on the server.</b>';if(pub)$('learnLbl').classList.add('hide')});
$('drop').onclick=()=>$('file').click();
['dragover','dragenter'].forEach(ev=>$('drop').addEventListener(ev,e=>{e.preventDefault();$('drop').classList.add('on')}));
['dragleave','drop'].forEach(ev=>$('drop').addEventListener(ev,e=>{e.preventDefault();$('drop').classList.remove('on')}));
$('drop').addEventListener('drop',e=>send(e.dataTransfer.files[0]));$('file').onchange=e=>send(e.target.files[0]);
let timer=null;
function setStage(t,p,busy){$('stmsg').textContent=t;const b=$('barfill');b.classList.toggle('busy',!!busy);if(!busy)b.style.width=p+'%'}
function fail(m){clearInterval(timer);$('stage').classList.add('hide');$('msg').textContent=m}
function send(f){if(!f)return;if(f.size>25*1024*1024&&pub){$('msg').textContent='File too large (max 25 MB).';return}
$('msg').textContent='';$('out').classList.add('hide');$('stage').classList.remove('hide');setStage('Uploading '+f.name+' ...',0);
const fd=new FormData();fd.append('file',f);const x=new XMLHttpRequest();x.open('POST','/api/upload');
x.upload.onprogress=e=>{if(e.lengthComputable)setStage('Uploading '+f.name+' ... '+Math.round(e.loaded/e.total*100)+'%',e.loaded/e.total*100)};
x.upload.onload=()=>{const t0=Date.now();const tick=()=>setStage('\u2714 Uploaded successfully. Coding your file, please wait ('+Math.round((Date.now()-t0)/1000)+'s) - large files can take a few minutes ...',100,true);tick();timer=setInterval(tick,1000)};
x.onload=()=>{clearInterval(timer);let j;try{j=JSON.parse(x.responseText)}catch(e){return fail('The server took too long or restarted. Please try again.')}
if(x.status!=200)return fail(j.error||'Something went wrong.');setStage('\u2714 Coding complete. Preparing your results ...',100);
setTimeout(()=>{$('stage').classList.add('hide');show(j);$('view').classList.remove('hide');$('out').scrollIntoView({behavior:'smooth'})},1500)};
x.onerror=()=>fail('Upload failed. Check your connection and try again.');x.send(fd);$('file').value=''}
function show(j){cur=j;$('out').classList.remove('hide');$('sRows').textContent=j.rows;$('sE').textContent=j.exact;$('sC').textContent=j.close+j.likely;$('sR').textContent=j.review;$('sK').textContent=j.kesic;
$('tcol').innerHTML=j.columns.map(c=>'<option'+(c==j.tcol?' selected':'')+'>'+c+'</option>').join('');
$('ccol').innerHTML='<option value="">(none - skip KeSic)</option>'+j.columns.map(c=>'<option'+(c==j.ccol?' selected':'')+'>'+c+'</option>').join('');
$('dx').href='/download/'+j.id;$('dc').href='/download/'+j.id+'?fmt=csv';render()}
$('remap').onclick=async()=>{const r=await fetch('/api/remap',{method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify({id:cur.id,tcol:$('tcol').value,ccol:$('ccol').value})});show(await r.json())};
$('learn').onchange=async e=>{const fd=new FormData();fd.append('file',e.target.files[0]);fd.append('tcol',cur.tcol);fd.append('ccol',cur.ccol);
const r=await fetch('/api/learn',{method:'POST',body:fd});const j=await r.json();alert(r.ok?'Learned '+j.titles+' titles and '+j.companies+' companies. Re-code to use them.':j.error)};
$('viewBtn').onclick=()=>$('view').classList.toggle('hide');$('flt').onchange=render;$('q').oninput=render;
const esc=s=>String(s??'').replace(/[&<>"]/g,c=>({'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;'}[c]));
function render(){if(!cur)return;const q=$('q').value.toLowerCase(),f=$('flt').value,cols=cur.cols.filter(c=>c!='needs_review');
let d=cur.data.filter(r=>(!f||(f=='r')==!!r.needs_review)&&(!q||JSON.stringify(r).toLowerCase().includes(q)));
$('cnt').textContent=d.length+' rows'+(cur.rows>2000?' (preview of first 2,000 - download has all '+cur.rows+')':'');
$('t').innerHTML='<tr>'+cols.map(c=>'<th class="'+(/^kes(co|ic) /.test(c)?'k':'')+'">'+esc(c)+'</th>').join('')+'</tr>'+
d.slice(0,500).map(r=>'<tr class="'+(r.needs_review?'r':'')+'">'+cols.map(c=>'<td class="'+(/^kes(co|ic) /.test(c)?'k':'')+'" title="'+esc(r[c])+'">'+esc(r[c])+'</td>').join('')+'</tr>').join('')}
</script></body></html>""".replace("CSSHERE", PAGE_CSS)

if __name__ == "__main__":
    threading.Timer(1.2, lambda: webbrowser.open("http://127.0.0.1:5000")).start()
    app.run(host="127.0.0.1", port=int(os.environ.get("PORT", 5000)))
